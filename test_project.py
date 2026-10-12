# Run all project checks with: python -B test_project.py
# Tests use temporary files and simulated AI; they do not read keys or call the API.

import csv
import importlib.util
import json
import os
import runpy
import shutil
import sys
import traceback
from contextlib import redirect_stdout
from copy import deepcopy
from datetime import date
from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import Mock, patch

import data_manager
import io_manager
import logic_manager



# CSV storage and audit logging


def make_result(**changes):
    results = {
        "final_score": 80.0,
        "skills_score": 75.0,
        "experience_score": 90.0,
        "education_score": 70.0,
        "ai_overall_score": 85.0,
        "confidence": "high",
        "status": "SHORTLIST",
        "fast_track": False,
        "missing_must_haves": [],
        "stale_skills": [],
        "skill_details": [
            {"name": "Python", "category": "must_have", "present": True,
             "last_used_year": 2026, "evidence": "6 years of Python"},
            {"name": "Mircosoft", "category": "preferred", "present": False,
             "last_used_year": None, "evidence": ""}  
        ],
        "summary": "Strong fit",
        "notes": ["Must-have bonus +5"],

    }
    results.update(changes)
    return results

def save_one(files, file_name="alice.pdf", job=None, **changes):
    candidate = data_manager.result_to_candidate(file_name, "resume/" + file_name, make_result(**changes))
    return data_manager.save_candidate_result(candidate, job,
        results_file=files["results_file"],
        log_file=files["log_file"])

def test_list_text_round_trip():
    items = ["Python", "SQL", "C++"]
    assert data_manager.text_to_list(data_manager.list_to_text(items)) == items
    assert data_manager.text_to_list("") == []


def test_csv_write_creates_and_replaces_complete_file():
    # CSV quoting, Unicode, and replacement must survive the temporary-file write.
    with TemporaryDirectory() as folder:
        path = Path(folder) / "results.csv"
        fields = ["candidate_id", "summary"]
        data_manager.write_csv(path, fields, [{"candidate_id": "C001", "summary": "Old"}])
        rows = [{"candidate_id": "C002", "summary": 'Python, SQL\nRésumé "evidence"'}]
        data_manager.write_csv(path, fields, rows)
        assert data_manager.read_csv(path, fields) == rows
        assert list(Path(folder).iterdir()) == [path]


def test_default_storage_uses_application_data_folder():
    # Exercise defaults in an isolated application launched from another folder.
    with TemporaryDirectory() as folder:
        application = Path(folder) / "application"
        application.mkdir()
        module_path = application / "data_manager.py"
        shutil.copyfile(data_manager.__file__, module_path)
        previous_directory = os.getcwd()
        try:
            os.chdir(folder)
            storage = runpy.run_path(str(module_path))
            data_folder = application / "data"
            assert storage["load_all_results"]() == []
            assert storage["load_all_jobs"]() == []
            assert not data_folder.exists()
            job_id = storage["save_job"]({"title": "Developer"}, {})
            candidate_id = storage["save_candidate_result"](
                {"file_name": "sample.pdf", "final_score": 80, "status": "SHORTLIST"},
                {"job_id": job_id, "title": "Developer"},
            )
            storage["save_results_json"]({"sample": "debug output"})
            storage["update_recruiter_decision"](candidate_id, "shortlisted")
            assert storage["load_all_results"]()[0]["recruiter_decision"] == "shortlisted"
            assert storage["load_all_jobs"]()[0]["job_id"] == job_id
            assert len(storage["load_log"]()) == 3
            assert json.loads((data_folder / "resume_results.json").read_text(encoding="utf-8")) == {
                "sample": "debug output",
            }
            assert sorted(path.name for path in data_folder.iterdir()) == [
                "audit_log.csv", "candidate_results.csv", "jobs.csv", "resume_results.json",
            ]
            assert list(Path(folder).iterdir()) == [application]
        finally:
            os.chdir(previous_directory)


def test_storage_folder_failures_reach_caller():
    # A file blocking the folder must produce a clear failure for CSV and JSON.
    with TemporaryDirectory() as folder:
        blocked_folder = Path(folder) / "data"
        blocked_folder.write_text("Existing file", encoding="utf-8")
        with redirect_stdout(StringIO()):
            try:
                data_manager.write_csv(blocked_folder / "results.csv", ["candidate_id"], [])
            except OSError:
                pass
            else:
                raise AssertionError("CSV folder creation failure should reach the caller")
            try:
                data_manager.save_results_json({}, blocked_folder / "debug.json")
            except ValueError as error:
                assert "Could not save results" in str(error)
            else:
                raise AssertionError("JSON folder creation failure should reach the caller")
        assert blocked_folder.read_text(encoding="utf-8") == "Existing file"


def test_failed_csv_write_preserves_original_file():
    # Fail after writing one temporary row; the original must remain untouched.
    with TemporaryDirectory() as folder:
        path = Path(folder) / "results.csv"
        fields = ["candidate_id"]
        data_manager.write_csv(path, fields, [{"candidate_id": "C001"}])
        original = path.read_bytes()
        try:
            data_manager.write_csv(path, fields, [
                {"candidate_id": "C002"},
                {"candidate_id": "C003", "unexpected_field": "Invalid"},
            ])
        except ValueError:
            pass
        else:
            raise AssertionError("Invalid CSV row should raise an error")
        assert path.read_bytes() == original
        assert list(Path(folder).iterdir()) == [path]


def test_failed_csv_replacement_preserves_decision_and_audit():
    # Simulate an Excel lock; a failed decision save must not add a success audit.
    with TemporaryDirectory() as folder:
        results_file = Path(folder) / "results.csv"
        log_file = Path(folder) / "audit.csv"
        candidate_id = data_manager.save_candidate_result(
            {"file_name": "sample.pdf", "final_score": 80, "status": "SHORTLIST"},
            results_file=results_file, log_file=log_file,
        )
        original_results = results_file.read_bytes()
        original_log = log_file.read_bytes()
        with patch.object(data_manager.os, "replace", side_effect=PermissionError("CSV locked")), redirect_stdout(StringIO()):
            try:
                data_manager.update_recruiter_decision(
                    candidate_id, "shortlisted", results_file, log_file,
                )
            except PermissionError:
                pass
            else:
                raise AssertionError("Replacement failure should reach the caller")
        assert results_file.read_bytes() == original_results
        assert log_file.read_bytes() == original_log
        assert sorted(path.name for path in Path(folder).iterdir()) == ["audit.csv", "results.csv"]

def test_text_to_number_bad_values():
    assert data_manager.text_to_number("92.5") == 92.5
    assert data_manager.text_to_number("7", as_int=True) == 7
    assert data_manager.text_to_number("") is None
    assert data_manager.text_to_number("abc") is None

def test_bool_to_text():
    assert data_manager.bool_to_text(True) == "yes"
    assert data_manager.bool_to_text(False) == "no"
    assert data_manager.bool_to_text("True") == "yes"

def test_log_action_adds_a_timestamped_line():
    with TemporaryDirectory() as folder:
        files = {
            "results_file": str(Path(folder) / "results.csv"),
            "log_file": str(Path(folder) / "log.csv"),
        }
        data_manager.log_action("TEST_ACTION", "some details", "C001", files["log_file"])
        log = data_manager.load_log(files["log_file"])
        assert len(log) == 1
        assert log[0]["action"] == "TEST_ACTION"
        assert log[0]["candidate_id"] == "C001"
        assert log[0]["timestamp"] != ""

def test_skills_are_split_into_present_and_missing():
    candidate = data_manager.result_to_candidate("a.pdf", "resume/a.pdf", make_result())
    ai_result = candidate["ai_result"]
    assert ai_result["skills_present"] == ["Python"]
    assert ai_result["skills_missing"] == ["Mircosoft"]
    assert ai_result["evidence"] == "Python: 6 years of Python"

def test_candidate_row_round_trip():
    candidate = data_manager.result_to_candidate("a.pdf", "resume/a.pdf", make_result(fast_track=True))
    candidate["candidate_id"] = "C001"
    row = data_manager.candidate_to_row(data_manager.flatten_candidate(candidate))
    loaded = data_manager.row_to_candidate(row)
    assert loaded["file_name"] == "a.pdf"
    assert loaded["final_score"] == 80.0
    assert loaded["fast_track"] is True
    assert loaded["skills_present"] == ["Python"]
    assert loaded["ai_result"]["confidence"] == "high"

def test_ids_count_up():
    with TemporaryDirectory() as folder:
        files = {
            "results_file": str(Path(folder) / "results.csv"),
            "log_file": str(Path(folder) / "log.csv"),
        }
        assert save_one(files, "a.pdf") == "C001"
        assert save_one(files, "b.pdf") == "C002"

def test_job_title_comes_from_job_description():
    with TemporaryDirectory() as folder:
        files = {
            "results_file": str(Path(folder) / "results.csv"),
            "log_file": str(Path(folder) / "log.csv"),
        }
        candidate_id = save_one(files, job={"title": "Data Engineer"})
        assert data_manager.get_candidate(candidate_id, files["results_file"])["job_title"] == "Data Engineer"

def test_invalid_status_is_rejected():
    with TemporaryDirectory() as folder:
        files = {
            "results_file": str(Path(folder) / "results.csv"),
            "log_file": str(Path(folder) / "log.csv"),
        }
        try:
            save_one(files, status="MAYBE")
        except ValueError:
            return
        raise AssertionError("Invalid status should be rejected")

def test_resume_text_and_sensitive_fields_are_not_saved():
    with TemporaryDirectory() as folder:
        files = {
            "results_file": str(Path(folder) / "results.csv"),
            "log_file": str(Path(folder) / "log.csv"),
        }
        candidate = data_manager.result_to_candidate("a.pdf", "resume/a.pdf", make_result())
        candidate["raw_text"] = "full resume text"
        candidate["age"] = 30
        candidate["gender"] = "F"
        candidate_id = data_manager.save_candidate_result(candidate, None,
            results_file=files["results_file"],
            log_file=files["log_file"])
        loaded = data_manager.get_candidate(candidate_id, files["results_file"])
        for key in ["raw_text", "age", "gender"]:
            assert key not in loaded
        with open(files["results_file"], encoding="utf-8") as file:
            assert "full resume text" not in file.read()

def test_notes_with_semicolons_survive_a_round_trip():
    with TemporaryDirectory() as folder:
        files = {
            "results_file": str(Path(folder) / "results.csv"),
            "log_file": str(Path(folder) / "log.csv"),
        }
        notes = ["AI labelled 'Docker' as preferred; job description says must_have", "Missing must-haves: SQL, ETL; score capped"]
        candidate_id = save_one(files, notes=notes)
        assert data_manager.get_candidate(candidate_id, files["results_file"])["notes"] == notes

def test_save_screening_record_success():
    with TemporaryDirectory() as folder:
        files = {
            "results_file": str(Path(folder) / "results.csv"),
            "log_file": str(Path(folder) / "log.csv"),
        }
        record = {"resume_name": "a.pdf", "status": "success", "error": None}
        resume = {"file_name": "a.pdf", "file_path": "resume/a.pdf"}
        candidate_id = data_manager.save_screening_record(
            record, resume, make_result(), {"title": "Software Engineer"},
            results_file=files["results_file"],
            log_file=files["log_file"])
        loaded = data_manager.get_candidate(candidate_id, files["results_file"])
        assert loaded["file_name"] == "a.pdf"
        assert loaded["file_path"] == "resume/a.pdf"
        assert loaded["job_title"] == "Software Engineer"
        assert loaded["final_score"] == 80.0
        assert loaded["status"] == "SHORTLIST"

def test_save_screening_record_failed_ai_call():
    with TemporaryDirectory() as folder:
        files = {
            "results_file": str(Path(folder) / "results.csv"),
            "log_file": str(Path(folder) / "log.csv"),
        }
        record = {"resume_name": "b.docx", "status": "error", "error": "Failed after 3 attempts: timeout", "screening_result": None}
        resume = {"file_name": "b.docx", "file_path": "resume/b.docx"}
        candidate_id = data_manager.save_screening_record(
            record, resume, None, {"title": "Software Engineer"},
            results_file=files["results_file"],
            log_file=files["log_file"])
        loaded = data_manager.get_candidate(candidate_id, files["results_file"])
        assert loaded["status"] == "RE_ENTRY"
        assert loaded["final_score"] is None
        assert loaded["file_path"] == "resume/b.docx"
        assert loaded["notes"] == ["Failed after 3 attempts: timeout"]

def test_save_all_results_uses_resume_paths():
    with TemporaryDirectory() as folder:
        files = {
            "results_file": str(Path(folder) / "results.csv"),
            "log_file": str(Path(folder) / "log.csv"),
        }
        results = {"a.pdf": make_result(), "b.pdf": make_result(final_score=50.0, status="SECONDARY_REVIEW")}
        resumes = [{"file_name": "a.pdf", "file_path": "folder/a.pdf"}]
        ids = data_manager.save_all_results(results, resumes, {"title": "Analyst"},
            results_file=files["results_file"],
            log_file=files["log_file"])
        assert ids == ["C001", "C002"]
        assert data_manager.get_candidate("C001", files["results_file"])["file_path"] == "folder/a.pdf"
        assert data_manager.get_candidate("C002", files["results_file"])["file_path"] == ""


# Scoring and storage integration


def test_scored_skills_and_evidence_survive_storage():
    # Score a sample response, save it, and verify skills and evidence reload.
    job_description = {
        "title": "Developer",
        "required_skills": ["Python", "SQL"],
        "preferred": [],
    }
    ai_response = {
        "skills": [
            {
                "name": "Python",
                "category": "must_have",
                "present": True,
                "last_used_year": date.today().year,
                "evidence": "Built Python tools",
            },
            {
                "name": "SQL",
                "category": "must_have",
                "present": False,
                "last_used_year": None,
                "evidence": "",
            },
        ],
        "experience_score": 80,
        "education_score": 80,
        "overall_score": 70,
        "confidence": "high",
        "summary": "Python experience; SQL not demonstrated.",
    }

    result = logic_manager.evaluate_candidate(ai_response, job_description)
    candidate = data_manager.result_to_candidate("sample.pdf", "sample.pdf", result)

    with TemporaryDirectory() as folder:
        results_file = str(Path(folder) / "results.csv")
        log_file = str(Path(folder) / "audit.csv")
        candidate_id = data_manager.save_candidate_result(
            candidate, job_description, results_file, log_file
        )
        loaded = data_manager.get_candidate(candidate_id, results_file)

    assert loaded["ai_result"]["skills_present"] == ["Python"]
    assert loaded["ai_result"]["skills_missing"] == ["SQL"]
    assert loaded["ai_result"]["evidence"] == "Python: Built Python tools"
    assert loaded["final_score"] == result["final_score"]
    assert loaded["missing_must_haves"] == ["SQL"]
    assert loaded["summary"] == ai_response["summary"]


# Screening orchestration


def test_main_scores_and_saves_when_debug_write_fails():
    # Keep a successful candidate and an unscored failure despite debug errors.
    ai_response = {
        "skills": [{
            "name": "Python", "category": "must_have", "present": True,
            "last_used_year": date.today().year, "evidence": "Built Python tools",
        }],
        "experience_score": 80, "education_score": 80,
        "overall_score": 80, "confidence": "high", "summary": "Python developer",
    }
    simulated_ai = SimpleNamespace(
        create_client=Mock(return_value=object()),
        process_resume_ai=Mock(side_effect=[ai_response, ValueError("Invalid API response")]),
    )
    # Replace the AI module before loading main: no .env loading or API access.
    module_spec = importlib.util.spec_from_file_location(
        "screening_flow_test", Path(__file__).with_name("main.py")
    )
    orchestrator = importlib.util.module_from_spec(module_spec)
    with patch.dict(sys.modules, {"ai_manager": simulated_ai}):
        module_spec.loader.exec_module(orchestrator)

    resumes = [
        {"id": "one", "file_name": "good.pdf", "file_path": "good.pdf", "text": "Built Python tools"},
        {"id": "two", "file_name": "failed.pdf", "file_path": "failed.pdf", "text": "Sample text"},
    ]
    job = {"title": "Developer", "required_skills": ["Python"], "preferred": []}

    with TemporaryDirectory() as folder:
        results_file = str(Path(folder) / "results.csv")
        log_file = str(Path(folder) / "audit.csv")
        jobs_file = str(Path(folder) / "jobs.csv")

        def save_record(record, resume, evaluation, job_description):
            # Use real storage with temporary output paths for this test.
            return data_manager.save_screening_record(
                record, resume, evaluation, job_description, results_file, log_file
            )

        with patch.multiple(
            orchestrator,
            collect_inputs=Mock(return_value=(resumes, job, {"min_qualifying_score": 60})),
            save_results_json=Mock(side_effect=ValueError("Debug folder is read-only")),
            save_screening_record=save_record,
            save_job=lambda description, preferences: data_manager.save_job(description, preferences, jobs_file, log_file),
            show_message=Mock(), show_error=Mock(), show_progress=Mock(),
            show_summary=Mock(), show_rankings=Mock(),
            get_main_menu_choice=Mock(side_effect=["1", "4"]),
        ):
            orchestrator.main()
            rows = orchestrator.show_rankings.call_args.args[0]
            assert len(rows) == 1
            assert rows[0]["file_name"] == "good.pdf"
            assert any("Warning:" in call.args[0] for call in orchestrator.show_message.call_args_list)
            orchestrator.show_error.assert_not_called()

        candidates = data_manager.load_all_results(results_file)
        audit = data_manager.load_log(log_file)

    assert len(candidates) == 2
    assert candidates[0]["skills_present"] == ["Python"]
    assert candidates[0]["evidence"] == "Python: Built Python tools"
    assert candidates[0]["recruiter_decision"] == "pending"
    assert candidates[1]["status"] == "RE_ENTRY"
    assert candidates[1]["final_score"] is None
    assert candidates[1]["notes"] == ["Invalid API response"]
    assert all(candidate["job_id"] == "J001" for candidate in candidates)
    assert len(audit) == 3
    assert audit[0]["action"] == "CREATE_JOB"
    assert all(entry["action"] == "SAVE_RESULT" for entry in audit[1:])


# Saved-results menu


def load_test_orchestrator():
    # Substitute the AI module before loading main so no keys or .env are read.
    simulated_ai = SimpleNamespace(create_client=Mock(), process_resume_ai=Mock())
    spec = importlib.util.spec_from_file_location(
        "saved_results_test", Path(__file__).with_name("main.py")
    )
    orchestrator = importlib.util.module_from_spec(spec)
    with patch.dict(sys.modules, {"ai_manager": simulated_ai}):
        spec.loader.exec_module(orchestrator)
    return orchestrator, simulated_ai

def test_menu_views_csv_without_changing_records():
    # View both scored and failed records, retry invalid selections, and leave CSV unchanged.
    orchestrator, simulated_ai = load_test_orchestrator()
    with TemporaryDirectory() as folder:
        results_file = str(Path(folder) / "results.csv")
        log_file = str(Path(folder) / "audit.csv")
        candidate = data_manager.result_to_candidate("good.pdf", "good.pdf", {
            "final_score": 85, "skills_score": 80, "experience_score": 90,
            "education_score": 85, "status": "SHORTLIST", "confidence": "high",
            "summary": "Relevant Python experience", "notes": ["Must-have bonus +5"],
            "skill_details": [{"name": "Python", "present": True, "evidence": "Built Python tools"}],
        })
        data_manager.save_candidate_result(candidate, {"title": "Developer"}, results_file, log_file)
        data_manager.save_failed_resume("failed.pdf", reason="API timeout", results_file=results_file, log_file=log_file)
        original_results = Path(results_file).read_bytes()
        original_audit = Path(log_file).read_bytes()
        output = StringIO()
        with patch.object(orchestrator, "load_all_results", side_effect=lambda: data_manager.load_all_results(results_file)), patch(
            "builtins.input", side_effect=["bad", "2", "1", "C404", " c001 ", "C002", "", "4"]
        ), patch.object(orchestrator, "load_all_jobs", return_value=[]), patch.object(
            orchestrator, "load_job_results", side_effect=lambda job_id: data_manager.load_job_results(job_id, results_file)
        ), redirect_stdout(output):
            orchestrator.main()
        assert Path(results_file).read_bytes() == original_results
        assert Path(log_file).read_bytes() == original_audit
    simulated_ai.create_client.assert_not_called()
    simulated_ai.process_resume_ai.assert_not_called()
    text = output.getvalue()
    for expected in ["Invalid choice", "Candidate ID not found", "Built Python tools",
                     "Relevant Python experience", "Not scored", "API timeout", "Recruiter decision: pending"]:
        assert expected in text

def test_empty_results_and_read_errors_return_to_menu():
    # A missing CSV is not created, and failed loading still allows normal exit.
    orchestrator, simulated_ai = load_test_orchestrator()
    with TemporaryDirectory() as folder:
        missing_file = str(Path(folder) / "missing.csv")
        output = StringIO()
        with patch.object(orchestrator, "load_all_results", side_effect=lambda: data_manager.load_all_results(missing_file)), patch(
            "builtins.input", side_effect=["2", "4"]
        ), patch.object(orchestrator, "load_all_jobs", return_value=[]), redirect_stdout(output):
            orchestrator.main()
        assert "No saved results yet" in output.getvalue()
        assert not Path(missing_file).exists()
    output = StringIO()
    with patch.object(orchestrator, "load_all_results", side_effect=PermissionError("File is locked")), patch(
        "builtins.input", side_effect=["2", "4"]
    ), redirect_stdout(output):
        orchestrator.main()
    assert "Could not load saved results" in output.getvalue()
    assert "Exiting program" in output.getvalue()
    simulated_ai.create_client.assert_not_called()


# Job storage and selection


def test_jobs_keep_identical_titles_separate():
    # Save distinct requirements under identical titles and view only the selected job.
    orchestrator, simulated_ai = load_test_orchestrator()
    with TemporaryDirectory() as folder:
        jobs_file = str(Path(folder) / "jobs.csv")
        results_file = str(Path(folder) / "results.csv")
        log_file = str(Path(folder) / "audit.csv")
        job_one = {"title": "Developer", "required_skills": ["Python"], "preferred": [],
                   "min_experience": 3, "education": "Computing degree"}
        job_two = dict(job_one, required_skills=["SQL"], min_experience=5)
        for job in (job_one, job_two):
            job["job_id"] = data_manager.save_job(job, {"min_qualifying_score": 70}, jobs_file, log_file)
            result = {"final_score": 80, "status": "SHORTLIST"}
            candidate = data_manager.result_to_candidate(job["job_id"] + ".pdf", "sample.pdf", result)
            data_manager.save_candidate_result(candidate, job, results_file, log_file)
        failed_record = {"status": "error", "error": "API timeout"}
        data_manager.save_screening_record(failed_record, {"file_name": "failed.pdf"},
                                          None, job_one, results_file, log_file)
        assert job_one["job_id"] == "J001"
        assert job_two["job_id"] == "J002"
        jobs = data_manager.load_all_jobs(jobs_file)
        assert jobs[0]["required_skills"] == ["Python"]
        assert jobs[1]["required_skills"] == ["SQL"]
        assert jobs[1]["min_experience"] == 5
        assert jobs[1]["min_qualifying_score"] == 70
        assert len(data_manager.load_job_results("J001", results_file)) == 2
        assert len(data_manager.load_job_results("J002", results_file)) == 1
        assert data_manager.load_job_results("J001", results_file)[1]["final_score"] is None
        snapshots = {path: Path(path).read_bytes() for path in (jobs_file, results_file, log_file)}
        output = StringIO()
        with patch.multiple(
            orchestrator,
            load_all_jobs=lambda: data_manager.load_all_jobs(jobs_file),
            load_all_results=lambda: data_manager.load_all_results(results_file),
            load_job_results=lambda job_id: data_manager.load_job_results(job_id, results_file),
        ), patch("builtins.input", side_effect=["2", "bad", "9", "2", "C001", "C002", "", "4"]), redirect_stdout(output):
            orchestrator.main()
        text = output.getvalue()
        assert "Required skills: SQL" in text
        assert "Minimum experience: 5.0 years" in text
        assert "Candidate ID not found" in text
        assert "J002.pdf" in text
        assert "J001.pdf" not in text
        for path, original in snapshots.items():
            assert Path(path).read_bytes() == original
    simulated_ai.create_client.assert_not_called()

def test_older_csv_rows_survive_new_job_column():
    # Simulate the old CSV header, then verify saving a new row preserves old values.
    with TemporaryDirectory() as folder:
        results_file = str(Path(folder) / "results.csv")
        log_file = str(Path(folder) / "audit.csv")
        old_fields = [field for field in data_manager.CANDIDATE_FIELDS if field != "job_id"]
        old_row = {field: "" for field in old_fields}
        old_row.update(candidate_id="C001", file_name="older.pdf", final_score="82.5",
                       status="SHORTLIST", recruiter_decision="shortlisted", notes="Existing notes")
        with open(results_file, "w", newline="", encoding="utf-8") as file:
            writer = csv.DictWriter(file, fieldnames=old_fields)
            writer.writeheader()
            writer.writerow(old_row)
        assert data_manager.load_job_results("", results_file)[0]["final_score"] == 82.5
        candidate = {"file_name": "new.pdf", "final_score": 80, "status": "SHORTLIST"}
        data_manager.save_candidate_result(candidate, {"title": "Developer", "job_id": "J001"}, results_file, log_file)
        old_candidate = data_manager.load_job_results("", results_file)[0]
        assert old_candidate["file_name"] == "older.pdf"
        assert old_candidate["recruiter_decision"] == "shortlisted"
        assert old_candidate["notes"] == ["Existing notes"]
        assert len(data_manager.load_job_results("J001", results_file)) == 1

def test_job_save_failure_stops_screening():
    # A failed job write must be reported before any API screening request.
    orchestrator, simulated_ai = load_test_orchestrator()
    with patch.multiple(
        orchestrator,
        collect_inputs=Mock(return_value=([{"id": "one"}], {"title": "Developer"}, {})),
        save_job=Mock(side_effect=PermissionError("Jobs file is locked")),
        show_error=Mock(),
    ):
        orchestrator.screen_new_batch()
        assert "Could not finish saving the job" in orchestrator.show_error.call_args.args[0]
    simulated_ai.process_resume_ai.assert_not_called()


# Skill evidence


def test_skill_evidence_flow():
    # Disable dotenv before loading the API validator; no client is created.
    with patch("dotenv.load_dotenv"):
        ai_functions = runpy.run_path(str(Path(__file__).with_name("ai_manager.py")))
    validate = ai_functions["validate_ai_result"]
    job = {"title": "Developer", "required_skills": ["Python", "SQL"], "preferred": []}
    response = {
        "skills": [
            {"name": "Python", "category": "must_have", "present": True,
             "last_used_year": date.today().year, "evidence": "Built Python scripts"},
            {"name": "SQL", "category": "must_have", "present": True,
             "last_used_year": date.today().year, "evidence": "Wrote SQL queries"},
        ],
        "experience_score": 100, "education_score": 100,
        "overall_score": 100, "confidence": "high",
    }
    supported = logic_manager.evaluate_candidate(validate(response), job)
    assert supported["final_score"] == 100

    for evidence in ["", "   ", None]:
        unsupported = deepcopy(response)
        unsupported["skills"][1]["evidence"] = evidence
        result = logic_manager.evaluate_candidate(validate(unsupported), job)
        assert result["skills_score"] == 50
        assert result["final_score"] == 75
        assert result["status"] != "RE_ENTRY"
        assert result["missing_must_haves"] == []
        assert result["fast_track"] is False
        assert not any("bonus" in note or "score capped" in note for note in result["notes"])

    del unsupported["skills"][1]["evidence"]
    assert logic_manager.evaluate_candidate(validate(unsupported), job)["final_score"] == 75

    missing = deepcopy(response)
    missing["skills"][1].update(present=False, evidence="")
    missing_result = logic_manager.evaluate_candidate(validate(missing), job)
    assert missing_result["final_score"] == 59
    assert missing_result["missing_must_haves"] == ["SQL"]

    malformed = deepcopy(response)
    malformed["skills"][1]["evidence"] = 123
    try:
        validate(malformed)
    except ValueError:
        pass
    else:
        raise AssertionError("Non-text evidence should be rejected")

    output = StringIO()
    with redirect_stdout(output):
        io_manager.show_rankings([{"file_name": "sample.pdf", "evaluation": result}])
    assert "Evidence: Built Python scripts" in output.getvalue()
    assert "SQL: CLAIMED WITHOUT EVIDENCE" in output.getvalue()
    assert "SQL: NOT DEMONSTRATED" not in output.getvalue()

    with TemporaryDirectory() as folder:
        results_file = str(Path(folder) / "results.csv")
        log_file = str(Path(folder) / "audit.csv")
        candidate = data_manager.result_to_candidate("sample.pdf", "sample.pdf", result)
        candidate_id = data_manager.save_candidate_result(candidate, job, results_file, log_file)
        loaded = data_manager.get_candidate(candidate_id, results_file)
    assert "SQL: claimed by AI without supporting evidence" in loaded["evidence"]
    assert loaded["final_score"] == 75
    assert any("no skill points" in note for note in loaded["notes"])


# Input validation and anonymization


def test_resume_validation_and_folder_reentry():
    # Reject invalid paths/files, retry unreadable text, and retain accepted resumes.
    with TemporaryDirectory() as folder:
        directory = Path(folder)
        (directory / "good.PDF").write_bytes(b"sample")
        (directory / "retry.docx").write_bytes(b"sample")
        (directory / "empty.pdf").touch()
        (directory / "wrong.txt").write_text("sample", encoding="utf-8")
        (directory / "directory.pdf").mkdir()
        attempts = {}

        def extract_sample_text(file_path):
            # Simulate a readable file and a failed extraction repaired on retry.
            name = Path(file_path).name
            attempts[name] = attempts.get(name, 0) + 1
            if name == "retry.docx" and attempts[name] == 1:
                return ""
            return "Python experience"

        messages = StringIO()
        with redirect_stdout(messages), patch("builtins.input", side_effect=[folder, folder, ""]), patch.object(
            io_manager, "extract_resume_text", side_effect=extract_sample_text
        ):
            resumes = io_manager.collect_resumes()

    assert [resume["file_name"] for resume in resumes] == ["good.PDF", "retry.docx"]
    assert attempts == {"good.PDF": 1, "retry.docx": 2}
    output = messages.getvalue()
    for reason in ["regular file", "Unsupported file type", "File is empty", "No readable text"]:
        assert reason in output

def test_job_inputs_reject_invalid_experience_and_clean_skills():
    # Re-prompt invalid experience and remove blank/duplicate skill entries.
    inputs = [
        "Developer", " Python, python, , SQL, SQL ",
        "abc", "-1", "nan", "inf", "-inf", "2.5",
        "Degree", "python, Docker, docker, , SQL",
    ]
    with redirect_stdout(StringIO()), patch("builtins.input", side_effect=inputs):
        job = io_manager.collect_job_description()
    assert job["required_skills"] == ["Python", "SQL"]
    assert job["preferred"] == ["Docker"]
    assert job["min_experience"] == 2.5
    assert io_manager.parse_skill_list(" , , ") == []

def test_anonymization_preserves_qualifications():
    # Remove synthetic personal details without changing skills or employment dates.
    text = (
        " Name: Alex Example\n"
        "ADDRESS: 12 Example Road\n"
        "Date of birth: 01/01/1990\n"
        "Age: 36\nGender: Male\nRace: Example\nMarital status: Single\n"
        "Contact me at alex@example.com or +65 8123 4567\n"
        "Alternative: (415) 555-0123\n"
        "International: +44 20 7946 0958\n"
        "Developer: 2019-2024\nEarlier role: 2010 - 2018\n"
        "Python 3.12, SQL; 6 years of experience\nBachelor of Computing"
    )
    cleaned = io_manager.anonymize_resume_text(text)
    for private_value in [
        "Alex Example", "12 Example Road", "01/01/1990", "Age:", "Gender:",
        "Race:", "Marital status:", "alex@example.com", "8123 4567",
        "415", "7946",
    ]:
        assert private_value not in cleaned
    for qualification in [
        "2019-2024", "2010 - 2018", "Python 3.12, SQL",
        "6 years of experience", "Bachelor of Computing",
    ]:
        assert qualification in cleaned
    assert "[EMAIL REMOVED]" in cleaned
    assert cleaned.count("[PHONE REMOVED]") == 3
    assert io_manager.anonymize_resume_text("+65 8123 4567 | 2020-2024") == "[PHONE REMOVED] | 2020-2024"
    assert io_manager.anonymize_resume_text("Name: Alex\nalex@example.com\n81234567") == ""

def test_prepared_resumes_use_anonymized_text():
    # Check the AI record uses processed text and rejects personal-details-only input.
    with TemporaryDirectory() as folder:
        resume_path = Path(folder) / "sample.pdf"
        resume_path.write_bytes(b"original document")
        with patch.object(io_manager, "extract_resume_text", return_value="Name: Alex\nPython experience"):
            resumes = io_manager.prepare_resumes_for_ai([str(resume_path)])
        assert resumes[0]["text"] == "Python experience"
        with patch.object(io_manager, "extract_resume_text", return_value="Name: Alex"), redirect_stdout(StringIO()) as output:
            assert io_manager.prepare_resumes_for_ai([str(resume_path)]) == []
        assert "No resume content remains" in output.getvalue()
        assert resume_path.read_bytes() == b"original document"

# Optional recruiter review

def test_optional_review_saves_decisions_and_preserves_scores():
    # Review can override recommendations, skip records, and survive interruption.
    orchestrator, simulated_ai = load_test_orchestrator()
    with TemporaryDirectory() as folder:
        results_file = str(Path(folder) / "results.csv")
        log_file = str(Path(folder) / "audit.csv")
        for index in range(1, 6):
            candidate = {
                "file_name": f"candidate_{index}.pdf", "final_score": 80,
                "status": "SHORTLIST", "recruiter_decision": "pending",
            }
            if index == 1:
                candidate.update(final_score=50, status="SECONDARY_REVIEW")
            elif index == 4:
                candidate.update(final_score=None, status="RE_ENTRY")
            elif index == 5:
                candidate["recruiter_decision"] = "rejected"
            data_manager.save_candidate_result(candidate, results_file=results_file, log_file=log_file)
        original_rows = data_manager.read_csv(results_file, data_manager.CANDIDATE_FIELDS)
        output = StringIO()
        with patch.multiple(
            orchestrator,
            load_all_results=lambda: data_manager.load_all_results(results_file),
            load_all_jobs=lambda: [],
            load_job_results=lambda job_id: data_manager.load_job_results(job_id, results_file),
            update_recruiter_decision=lambda candidate_id, decision: data_manager.update_recruiter_decision(
                candidate_id, decision, results_file, log_file
            ),
        ), patch("builtins.input", side_effect=[
            "3", "1", "bad", "1", "2", "3", "4", "C001", "3", "3", "4",
        ]), redirect_stdout(output):
            orchestrator.main()
        candidates = data_manager.load_all_results(results_file)
        assert [candidate["recruiter_decision"] for candidate in candidates] == [
            "shortlisted", "rejected", "pending", "pending", "rejected",
        ]
        updated_rows = data_manager.read_csv(results_file, data_manager.CANDIDATE_FIELDS)
        for original, updated in zip(original_rows, updated_rows):
            for field in data_manager.CANDIDATE_FIELDS:
                if field not in ("recruiter_decision", "updated_at"):
                    assert updated[field] == original[field]
        decisions = [entry for entry in data_manager.load_log(log_file) if entry["action"] == "RECRUITER_DECISION"]
        assert [entry["candidate_id"] for entry in decisions] == ["C001", "C002"]
        assert "Invalid choice" in output.getvalue()
        assert "needs successful screening" in output.getvalue()
        assert "Recruiter decision: shortlisted" in output.getvalue()
        # A decision saved before interruption must remain on disk.
        with patch.multiple(
            orchestrator,
            load_all_results=lambda: data_manager.load_all_results(results_file),
            load_all_jobs=lambda: [],
            load_job_results=lambda job_id: data_manager.load_job_results(job_id, results_file),
            update_recruiter_decision=lambda candidate_id, decision: data_manager.update_recruiter_decision(
                candidate_id, decision, results_file, log_file
            ),
        ), patch("builtins.input", side_effect=["3", "1", "3", "3", "1", EOFError()]), redirect_stdout(StringIO()):
            orchestrator.main()
        assert data_manager.get_candidate("C003", results_file)["recruiter_decision"] == "shortlisted"
        assert data_manager.get_candidate("C004", results_file)["recruiter_decision"] == "pending"
    simulated_ai.create_client.assert_not_called()
    simulated_ai.process_resume_ai.assert_not_called()


def test_review_reports_failed_or_missing_candidate_updates():
    # Failed saves must not report success or change the displayed decision.
    orchestrator, simulated_ai = load_test_orchestrator()
    for updater, expected in [
        (Mock(side_effect=PermissionError("CSV locked")), "Could not finish saving the decision"),
        (Mock(return_value=False), "Candidate no longer exists"),
    ]:
        candidate = data_manager.row_to_candidate({
            "candidate_id": "C001", "file_name": "sample.pdf",
            "final_score": "80", "status": "SHORTLIST", "recruiter_decision": "pending",
        })
        output = StringIO()
        with patch.multiple(
            orchestrator,
            load_all_results=lambda: [candidate], load_all_jobs=lambda: [],
            load_job_results=lambda job_id: [candidate], update_recruiter_decision=updater,
        ), patch("builtins.input", side_effect=["3", "1", "1", "4"]), redirect_stdout(output):
            orchestrator.main()
        assert expected in output.getvalue()
        assert "Decision saved" not in output.getvalue()
        assert candidate["recruiter_decision"] == "pending"
    simulated_ai.create_client.assert_not_called()


def test_sequential_review_returns_to_paused_candidate_after_jump():
    # A detour changes the target's decision without losing the sequential position.
    orchestrator, simulated_ai = load_test_orchestrator()
    with TemporaryDirectory() as folder:
        results_file = str(Path(folder) / "results.csv")
        log_file = str(Path(folder) / "audit.csv")
        for index in range(1, 4):
            data_manager.save_candidate_result(
                {"file_name": f"candidate_{index}.pdf", "final_score": 80, "status": "SHORTLIST"},
                results_file=results_file, log_file=log_file,
            )
        output = StringIO()
        with patch.multiple(
            orchestrator,
            load_all_results=lambda: data_manager.load_all_results(results_file),
            load_all_jobs=lambda: [],
            load_job_results=lambda job_id: data_manager.load_job_results(job_id, results_file),
            update_recruiter_decision=lambda candidate_id, decision: data_manager.update_recruiter_decision(
                candidate_id, decision, results_file, log_file
            ),
        ), patch("builtins.input", side_effect=[
            "3", "1", "3", "4", "C999", "c003", "1", "2", "3", "4",
        ]), redirect_stdout(output):
            orchestrator.main()
        displayed_ids = [
            line.split(": ", 1)[1] for line in output.getvalue().splitlines()
            if line.startswith("Candidate ID: ")
        ]
        assert displayed_ids == ["C001", "C002", "C003", "C002", "C003"]
        assert "Candidate ID not found" in output.getvalue()
        assert "Returning to C002" in output.getvalue()
        assert "Sequential review complete" in output.getvalue()
        candidates = data_manager.load_all_results(results_file)
        assert [candidate["recruiter_decision"] for candidate in candidates] == [
            "pending", "rejected", "shortlisted",
        ]
        decisions = [entry for entry in data_manager.load_log(log_file) if entry["action"] == "RECRUITER_DECISION"]
        assert [entry["candidate_id"] for entry in decisions] == ["C003", "C002"]
    simulated_ai.create_client.assert_not_called()


def test_cancelled_jump_and_early_exit_keep_decisions():
    # Cancelling a jump returns to the same candidate, and exiting does not save a decision.
    orchestrator, simulated_ai = load_test_orchestrator()
    candidates = [data_manager.row_to_candidate({
        "candidate_id": "C001", "file_name": "sample.pdf", "final_score": "80",
        "status": "SHORTLIST", "recruiter_decision": "pending",
    })]
    output = StringIO()
    updater = Mock()
    with patch.multiple(orchestrator, update_recruiter_decision=updater), patch(
        "builtins.input", side_effect=["4", "", "5"]
    ), redirect_stdout(output):
        orchestrator.review_saved_candidates(candidates)
    assert output.getvalue().count("Candidate ID: C001") == 2
    assert candidates[0]["recruiter_decision"] == "pending"
    updater.assert_not_called()
    simulated_ai.create_client.assert_not_called()


# Broader application checks and regression cases found during review


def load_test_ai():
    # Load real API functions with dotenv disabled; all requests use fake clients.
    with patch("dotenv.load_dotenv"):
        return runpy.run_path(str(Path(__file__).with_name("ai_manager.py")))


def make_ai_response():
    # Return a synthetic response with current, evidenced skills.
    return {
        "skills": [{"name": "Python", "category": "must_have", "present": True,
                    "last_used_year": date.today().year, "evidence": "Built Python tools"}],
        "experience_score": 80, "education_score": 80, "overall_score": 85,
        "confidence": "high", "summary": "Relevant Python experience",
    }


def wrap_api_response(payload):
    # Match the SDK response structure without making a network request.
    return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(payload)))])


def test_api_error_mapping_and_success():
    ai = load_test_ai()
    sdk = ai["openai"]
    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=Mock())))
    request = Mock()
    cases = [
        (sdk.APITimeoutError(request=request), ConnectionError, "timed out"),
        (sdk.APIConnectionError(request=request), ConnectionError, "connect"),
    ]
    for error_type, status, expected_type, message in [
        (sdk.RateLimitError, 429, ConnectionError, "rate limit"),
        (sdk.InternalServerError, 500, ConnectionError, "unavailable"),
        (sdk.AuthenticationError, 401, ValueError, "key was rejected"),
        (sdk.PermissionDeniedError, 403, ValueError, "key was rejected"),
        (sdk.NotFoundError, 404, ValueError, "model was not found"),
        (sdk.BadRequestError, 400, ValueError, "rejected the request"),
    ]:
        response = Mock(status_code=status, request=request, headers={})
        cases.append((error_type("Synthetic failure", response=response, body=None), expected_type, message))
    for error, expected_type, message in cases:
        client.chat.completions.create.side_effect = error
        try:
            ai["process_resume_ai"](client, {"required_skills": ["Python"]}, "Built Python tools")
        except expected_type as caught:
            assert message in str(caught)
        else:
            raise AssertionError(f"API failure was not reported: {type(error).__name__}")
    client.chat.completions.create.side_effect = None
    client.chat.completions.create.return_value = wrap_api_response(make_ai_response())
    assert ai["process_resume_ai"](client, {"title": "Developer"}, "Built Python tools") == make_ai_response()


def test_missing_key_and_existing_ai_validation():
    ai = load_test_ai()
    # Replace the environment access completely so the real key is never read.
    with patch.dict(ai["load_api_key"].__globals__, {"os": SimpleNamespace(environ={})}):
        try:
            ai["load_api_key"]()
        except ValueError as error:
            assert "not found" in str(error)
        else:
            raise AssertionError("Missing key should stop client creation")
    invalid = [[], {}, {**make_ai_response(), "skills": "Python"}]
    for score in [True, "80", -1, 101, float("nan"), float("inf")]:
        invalid.append({**make_ai_response(), "experience_score": score})
    invalid.append({**make_ai_response(), "confidence": "unknown"})
    for payload in invalid:
        try:
            ai["parse_response"](wrap_api_response(payload))
        except ValueError:
            pass
        else:
            raise AssertionError(f"Invalid response accepted: {payload}")
    for content in [None, "", "not JSON", "```json\n{}\n```"]:
        response = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=content))])
        try:
            ai["parse_response"](response)
        except ValueError:
            pass
        else:
            raise AssertionError("Empty or invalid JSON must be reported")


def test_unexpected_api_structure_is_reported_as_value_error():
    # Regression: malformed response containers must not crash the whole batch.
    ai = load_test_ai()
    for response in [SimpleNamespace(choices=[]), SimpleNamespace(choices=[None])]:
        try:
            ai["parse_response"](response)
        except ValueError:
            pass
        else:
            raise AssertionError("Malformed response should raise ValueError")


def test_real_document_readers_and_corrupt_files():
    from docx import Document
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    with TemporaryDirectory() as folder:
        directory = Path(folder)
        document = Document()
        document.add_paragraph("Name: Synthetic Candidate")
        document.add_paragraph("Built Python tools")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "SQL"
        table.cell(0, 1).text = "Database work"
        document.save(directory / "resume.docx")
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 200 Td (Built Python tools) Tj ET")
        page[NameObject("/Contents")] = stream
        writer.write(directory / "resume.pdf")
        blank = PdfWriter()
        blank.add_blank_page(width=300, height=300)
        blank.write(directory / "blank.pdf")
        encrypted = PdfWriter()
        encrypted.add_blank_page(width=300, height=300)
        encrypted.encrypt("synthetic-password")
        encrypted.write(directory / "encrypted.pdf")
        for name in ["corrupt.pdf", "corrupt.docx"]:
            (directory / name).write_bytes(b"This is not a valid document")
        output = StringIO()
        with redirect_stdout(output):
            resumes = io_manager.prepare_resumes_for_ai([str(path) for path in directory.iterdir()])
        assert sorted(resume["file_name"] for resume in resumes) == ["resume.docx", "resume.pdf"]
        assert all("Built Python tools" in resume["text"] for resume in resumes)
        assert all("Synthetic Candidate" not in resume["text"] for resume in resumes)
        assert "SQL | Database work" in next(r["text"] for r in resumes if r["file_name"].endswith("docx"))
        assert "Error reading" in output.getvalue()
        assert "No readable text" in output.getvalue()


def test_input_boundaries_and_folder_errors():
    with redirect_stdout(StringIO()):
        for value, expected in [("", 60), ("0", 0), ("100", 100)]:
            with patch("builtins.input", side_effect=["bad", "-1", "101", "nan", "inf", value]):
                assert io_manager.collect_screening_preferences()["min_qualifying_score"] == expected
        with patch.object(io_manager.os, "listdir", side_effect=PermissionError("Folder locked")):
            assert io_manager.collect_resume_batch("synthetic") == []
        with patch.object(io_manager, "collect_resumes", return_value=[]), patch.object(io_manager, "collect_job_description") as job:
            assert io_manager.collect_inputs() == ([], {}, {})
            job.assert_not_called()
        with patch("builtins.input", side_effect=["bad", "YES"]):
            assert io_manager.question_api_retry(1) is True
        with patch("builtins.input", return_value="no"):
            assert io_manager.question_api_retry(1) is False


def test_scoring_recency_bonus_threshold_and_fast_track():
    job = {"required_skills": ["Python"], "preferred": []}
    response = make_ai_response()
    response.update(experience_score=100, education_score=100, overall_score=100)
    result = logic_manager.evaluate_candidate(response, job, {"min_qualifying_score": 100})
    assert result["final_score"] == 100 and result["fast_track"]
    for age, expected in [(4, 100), (5, 50), (6, 50)]:
        response["skills"][0]["last_used_year"] = date.today().year - age
        result = logic_manager.evaluate_candidate(response, job)
        assert result["skills_score"] == expected
        assert result["stale_skills"] == ([] if age == 4 else ["Python"])
    response["skills"][0]["last_used_year"] = None
    result = logic_manager.evaluate_candidate(response, job)
    assert result["skills_score"] == 100
    assert any("dates unknown" in note for note in result["notes"])
    response.update(experience_score=80, education_score=80)
    assert logic_manager.evaluate_candidate(response, job)["final_score"] == 95
    assert logic_manager.evaluate_candidate(response, job, {"min_qualifying_score": 95})["status"] == "SHORTLIST"
    assert logic_manager.evaluate_candidate(response, job, {"min_qualifying_score": 96})["status"] == "SECONDARY_REVIEW"
    response["confidence"] = "low"
    assert not logic_manager.evaluate_candidate(response, job)["fast_track"]
    response["confidence"] = "high"
    response["education_score"] = 79
    assert not logic_manager.evaluate_candidate(response, job)["fast_track"]


def test_preferred_skills_missing_skills_and_ranking():
    response = make_ai_response()
    job = {"required_skills": ["Python"], "preferred": ["SQL"]}
    response["skills"][0]["category"] = "preferred"
    result = logic_manager.evaluate_candidate(response, job)
    assert result["skills_score"] == 66.7
    assert result["missing_must_haves"] == []
    assert any("job description says must_have" in note for note in result["notes"])
    job["required_skills"].append("Java")
    result = logic_manager.evaluate_candidate(response, job)
    assert result["missing_must_haves"] == ["Java"]
    assert result["final_score"] <= 59
    evaluations = {
        "low": {"fast_track": False, "final_score": 40},
        "high": {"fast_track": False, "final_score": 99},
        "fast": {"fast_track": True, "final_score": 95},
    }
    assert logic_manager.rank_candidates(evaluations) == ["fast", "high", "low"]


def test_full_screen_retry_view_and_review_flow():
    # Use real DOCX input, API parsing, scoring, CSV saving, and menu navigation.
    from docx import Document

    orchestrator, _ = load_test_orchestrator()
    ai = load_test_ai()
    request = Mock()
    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=Mock(side_effect=[
        ai["openai"].APITimeoutError(request=request),
        wrap_api_response(make_ai_response()),
        wrap_api_response(make_ai_response()),
    ]))))
    with TemporaryDirectory() as folder:
        directory = Path(folder)
        resumes = directory / "resumes"
        resumes.mkdir()
        for name in ["first.docx", "second.docx"]:
            document = Document()
            document.add_paragraph("Name: Synthetic Candidate")
            document.add_paragraph("Built Python tools")
            document.save(resumes / name)
        results = directory / "data" / "results.csv"
        audit = directory / "data" / "audit.csv"
        jobs = directory / "data" / "jobs.csv"
        with patch.multiple(
            orchestrator,
            create_client=lambda: client, process_resume_ai=ai["process_resume_ai"],
            DEBUG_JSON_PATH=str(directory / "data" / "debug.json"),
            save_job=lambda job, prefs: data_manager.save_job(job, prefs, jobs, audit),
            save_screening_record=lambda record, resume, score, job: data_manager.save_screening_record(record, resume, score, job, results, audit),
            load_all_results=lambda: data_manager.load_all_results(results),
            load_all_jobs=lambda: data_manager.load_all_jobs(jobs),
            load_job_results=lambda job_id: data_manager.load_job_results(job_id, results),
            update_recruiter_decision=lambda candidate_id, decision: data_manager.update_recruiter_decision(candidate_id, decision, results, audit),
        ), patch("builtins.input", side_effect=[
            "1", str(resumes), "", "Developer", "Python", "2", "Degree", "", "", "yes",
            "2", "1", "C001", "", "3", "1", "1", "3", "4",
        ]), redirect_stdout(StringIO()) as output:
            orchestrator.main()
        candidates = data_manager.load_all_results(results)
        assert len(candidates) == 2
        assert all(candidate["final_score"] == 95 for candidate in candidates)
        assert [candidate["recruiter_decision"] for candidate in candidates] == ["shortlisted", "pending"]
        assert "Retrying" in output.getvalue() and "Sequential review complete" in output.getvalue()
        assert client.chat.completions.create.call_count == 3
        for call in client.chat.completions.create.call_args_list:
            assert "Synthetic Candidate" not in call.kwargs["messages"][1]["content"]


def test_retries_only_repeat_temporary_failures():
    orchestrator, ai = load_test_orchestrator()
    resumes = {str(index): {"id": str(index), "file_name": f"{index}.pdf", "text": "Python"} for index in range(3)}
    ai.process_resume_ai.side_effect = [ConnectionError("Timeout"), ValueError("Invalid response"), make_ai_response(), ConnectionError("Timeout again"), make_ai_response()]
    results = {}
    orchestrator.screen_all_resumes(object(), resumes, {}, results)
    assert orchestrator.get_failed_ids(results) == ["0", "1"]
    assert orchestrator.retry_failed(object(), resumes, {}, results) == 1
    assert orchestrator.get_retryable_ids(results) == ["0"]
    assert orchestrator.retry_failed(object(), resumes, {}, results) == 1
    assert orchestrator.get_failed_ids(results) == ["1"]
    assert results["1"]["screening_result"] is None
    assert ai.process_resume_ai.call_count == 5


def test_startup_empty_input_and_menu_interruptions():
    orchestrator, ai = load_test_orchestrator()
    with patch.multiple(orchestrator, collect_inputs=Mock(), show_error=Mock()):
        ai.create_client.side_effect = ValueError("Missing synthetic key")
        orchestrator.screen_new_batch()
        orchestrator.collect_inputs.assert_not_called()
        assert "Missing synthetic key" in orchestrator.show_error.call_args.args[0]
    ai.create_client.side_effect = None
    with patch.multiple(orchestrator, collect_inputs=lambda: ([], {}, {}), save_job=Mock()), redirect_stdout(StringIO()):
        orchestrator.screen_new_batch()
        orchestrator.save_job.assert_not_called()
    for interruption in [EOFError(), KeyboardInterrupt()]:
        with patch("builtins.input", side_effect=interruption), redirect_stdout(StringIO()) as output:
            orchestrator.main()
        assert "Exiting" in output.getvalue()
    ai.process_resume_ai.assert_not_called()


def test_declined_retry_stores_failure_without_score():
    orchestrator, ai = load_test_orchestrator()
    ai.process_resume_ai.side_effect = ConnectionError("Synthetic timeout")
    resume = {"id": "one", "file_name": "sample.pdf", "text": "Python"}
    with TemporaryDirectory() as folder:
        results, audit, jobs = [Path(folder) / name for name in ["results.csv", "audit.csv", "jobs.csv"]]
        with patch.multiple(
            orchestrator,
            collect_inputs=lambda: ([resume], {"title": "Developer", "required_skills": ["Python"], "preferred": []}, {}),
            DEBUG_JSON_PATH=str(Path(folder) / "debug.json"),
            save_job=lambda job, prefs: data_manager.save_job(job, prefs, jobs, audit),
            save_screening_record=lambda record, item, score, job: data_manager.save_screening_record(record, item, score, job, results, audit),
        ), patch("builtins.input", return_value="no"), redirect_stdout(StringIO()):
            orchestrator.screen_new_batch()
        saved = data_manager.load_all_results(results)
        assert len(saved) == 1
        assert saved[0]["status"] == "RE_ENTRY" and saved[0]["final_score"] is None
        assert "Synthetic timeout" in saved[0]["notes"]
        assert ai.process_resume_ai.call_count == 1


def test_candidate_save_error_does_not_stop_remaining_saves():
    orchestrator, ai = load_test_orchestrator()
    ai.process_resume_ai.return_value = make_ai_response()
    resumes = [{"id": str(i), "file_name": f"{i}.pdf", "text": "Python"} for i in range(2)]
    saver = Mock(side_effect=[PermissionError("CSV locked"), "C001"])
    with patch.multiple(
        orchestrator,
        collect_inputs=lambda: (resumes, {"title": "Developer", "required_skills": ["Python"], "preferred": []}, {}),
        save_job=lambda job, prefs: "J001", save_results_json=Mock(),
        save_screening_record=saver,
    ), redirect_stdout(StringIO()) as output:
        orchestrator.screen_new_batch()
    assert saver.call_count == 2
    assert "CSV locked" in output.getvalue()
    assert "Saved 1/2" in output.getvalue()


def test_saved_jobs_without_candidates_and_missing_jobs():
    orchestrator, ai = load_test_orchestrator()
    job = {"job_id": "J001", "title": "Developer", "required_skills": ["Python"],
           "preferred": [], "min_experience": 0, "min_qualifying_score": 60}
    with patch.multiple(orchestrator, load_all_results=lambda: [], load_all_jobs=lambda: [job],
                        load_job_results=lambda job_id: []), patch("builtins.input", return_value="1"), redirect_stdout(StringIO()) as output:
        orchestrator.view_saved_results()
    assert "No candidates have been saved" in output.getvalue()
    candidate = data_manager.row_to_candidate({"candidate_id": "C001", "job_id": "J001"})
    with patch.multiple(orchestrator, load_all_results=lambda: [candidate], load_all_jobs=lambda: []), redirect_stdout(StringIO()) as output:
        orchestrator.view_saved_results()
    assert "no corresponding saved jobs" in output.getvalue()
    ai.create_client.assert_not_called()


def test_syntax_of_application_modules():
    # Compile source in memory to avoid creating bytecode or importing the API module.
    for name in ["main.py", "io_manager.py", "ai_manager.py", "logic_manager.py", "data_manager.py"]:
        path = Path(__file__).with_name(name)
        compile(path.read_text(encoding="utf-8"), str(path), "exec")


def test_invalid_logic_result_is_saved_without_score():
    # Regression: an empty skill list passes API validation but cannot be scored.
    ai = load_test_ai()
    response = ai["validate_ai_result"]({**make_ai_response(), "skills": []})
    job = {"required_skills": ["Python"], "preferred": []}
    evaluation = logic_manager.evaluate_candidate(response, job)
    assert evaluation["status"] == "RE_ENTRY"
    with TemporaryDirectory() as folder:
        results, audit = Path(folder) / "results.csv", Path(folder) / "audit.csv"
        data_manager.save_screening_record(
            {"status": "success"}, {"file_name": "sample.pdf"}, evaluation, job, results, audit,
        )
        saved = data_manager.load_all_results(results)[0]
        assert saved["final_score"] is None, "Unscorable AI output was stored as a zero score"


def test_interrupted_batch_preserves_completed_candidate():
    # Regression: a later interruption should not lose an already-screened candidate.
    orchestrator, ai = load_test_orchestrator()
    ai.process_resume_ai.side_effect = [make_ai_response(), KeyboardInterrupt()]
    resumes = [{"id": str(i), "file_name": f"{i}.pdf", "text": "Built Python tools"} for i in range(2)]
    with TemporaryDirectory() as folder:
        results, audit, jobs = [Path(folder) / name for name in ["results.csv", "audit.csv", "jobs.csv"]]
        with patch.multiple(
            orchestrator,
            collect_inputs=lambda: (resumes, {"title": "Developer", "required_skills": ["Python"], "preferred": []}, {}),
            DEBUG_JSON_PATH=str(Path(folder) / "debug.json"),
            save_job=lambda job, prefs: data_manager.save_job(job, prefs, jobs, audit),
            save_screening_record=lambda record, resume, score, job: data_manager.save_screening_record(record, resume, score, job, results, audit),
        ), redirect_stdout(StringIO()):
            orchestrator.screen_new_batch()
        assert len(data_manager.load_all_results(results)) == 1, "Completed candidate lost when next request was interrupted"


def test_debug_write_failure_cleans_temporary_file():
    with TemporaryDirectory() as folder:
        path = Path(folder) / "debug.json"
        data_manager.save_results_json({"previous": True}, path)
        original = path.read_bytes()
        try:
            data_manager.save_results_json({"invalid": object()}, path)
        except ValueError:
            pass
        else:
            raise AssertionError("Unserializable debug data should fail")
        assert path.read_bytes() == original
        assert list(Path(folder).iterdir()) == [path], "Failed debug save left a temporary file"


def test_invalid_candidate_csv_has_clear_error():
    # Regression: a missing candidate ID column should be reported before updates.
    with TemporaryDirectory() as folder:
        path = Path(folder) / "results.csv"
        path.write_text("file_name,final_score\nsample.pdf,80\n", encoding="utf-8")
        try:
            data_manager.generate_candidate_id(path)
        except ValueError:
            pass
        else:
            raise AssertionError("Missing candidate ID column should be reported")


def test_audit_failure_reports_partial_save():
    # Establish the current boundary: saving the CSV and audit are separate steps.
    with TemporaryDirectory() as folder:
        results = Path(folder) / "results.csv"
        with patch.object(data_manager, "log_action", side_effect=PermissionError("Audit locked")):
            try:
                data_manager.save_candidate_result(
                    {"file_name": "sample.pdf", "status": "SHORTLIST", "final_score": 80},
                    results_file=results, log_file=Path(folder) / "audit.csv",
                )
            except PermissionError:
                pass
            else:
                raise AssertionError("Audit failure must be reported")
        assert len(data_manager.load_all_results(results)) == 1


def test_json_loader_errors_and_saved_search_helpers():
    with TemporaryDirectory() as folder:
        path = Path(folder) / "debug.json"
        for payload in [None, "not JSON", "[]", '{"bad": {}}']:
            if payload is not None:
                path.write_text(payload, encoding="utf-8")
            try:
                logic_manager.load_results(path)
            except ValueError:
                pass
            else:
                raise AssertionError("Invalid saved JSON should be rejected")
        results, audit = Path(folder) / "results.csv", Path(folder) / "audit.csv"
        data_manager.save_candidate_result(
            {"file_name": "good.pdf", "status": "SHORTLIST", "final_score": 95,
             "fast_track": True, "summary": "Python", "job_title": "Developer"},
            results_file=results, log_file=audit,
        )
        data_manager.save_failed_resume("failed.pdf", results_file=results, log_file=audit)
        matches = data_manager.search_candidates(min_score=90, max_score=100, fast_track=True,
            status="SHORTLIST", decision="pending", job_title="developer", keyword="python", results_file=results)
        assert len(matches) == 1 and matches[0]["file_name"] == "good.pdf"
        assert data_manager.count_by_status(results) == {"SHORTLIST": 1, "SECONDARY_REVIEW": 0, "RE_ENTRY": 1}


def run_tests():
    # Run every check, report failures, and return an exit code for the terminal.
    tests = [
        test_list_text_round_trip,
        test_csv_write_creates_and_replaces_complete_file,
        test_default_storage_uses_application_data_folder,
        test_storage_folder_failures_reach_caller,
        test_failed_csv_write_preserves_original_file,
        test_failed_csv_replacement_preserves_decision_and_audit,
        test_text_to_number_bad_values,
        test_bool_to_text,
        test_log_action_adds_a_timestamped_line,
        test_skills_are_split_into_present_and_missing,
        test_candidate_row_round_trip,
        test_ids_count_up,
        test_job_title_comes_from_job_description,
        test_invalid_status_is_rejected,
        test_resume_text_and_sensitive_fields_are_not_saved,
        test_notes_with_semicolons_survive_a_round_trip,
        test_save_screening_record_success,
        test_save_screening_record_failed_ai_call,
        test_save_all_results_uses_resume_paths,
        test_scored_skills_and_evidence_survive_storage,
        test_main_scores_and_saves_when_debug_write_fails,
        test_menu_views_csv_without_changing_records,
        test_empty_results_and_read_errors_return_to_menu,
        test_jobs_keep_identical_titles_separate,
        test_older_csv_rows_survive_new_job_column,
        test_job_save_failure_stops_screening,
        test_skill_evidence_flow,
        test_resume_validation_and_folder_reentry,
        test_job_inputs_reject_invalid_experience_and_clean_skills,
        test_anonymization_preserves_qualifications,
        test_prepared_resumes_use_anonymized_text,
        test_optional_review_saves_decisions_and_preserves_scores,
        test_review_reports_failed_or_missing_candidate_updates,
        test_sequential_review_returns_to_paused_candidate_after_jump,
        test_cancelled_jump_and_early_exit_keep_decisions,
        test_api_error_mapping_and_success,
        test_missing_key_and_existing_ai_validation,
        test_unexpected_api_structure_is_reported_as_value_error,
        test_real_document_readers_and_corrupt_files,
        test_input_boundaries_and_folder_errors,
        test_scoring_recency_bonus_threshold_and_fast_track,
        test_preferred_skills_missing_skills_and_ranking,
        test_full_screen_retry_view_and_review_flow,
        test_retries_only_repeat_temporary_failures,
        test_startup_empty_input_and_menu_interruptions,
        test_declined_retry_stores_failure_without_score,
        test_candidate_save_error_does_not_stop_remaining_saves,
        test_saved_jobs_without_candidates_and_missing_jobs,
        test_syntax_of_application_modules,
        test_invalid_logic_result_is_saved_without_score,
        test_interrupted_batch_preserves_completed_candidate,
        test_debug_write_failure_cleans_temporary_file,
        test_invalid_candidate_csv_has_clear_error,
        test_audit_failure_reports_partial_save,
        test_json_loader_errors_and_saved_search_helpers,
    ]
    failures = 0
    for test in tests:
        try:
            test()
        except Exception:
            # Continue after a failed check, but display its complete error.
            failures += 1
            print(f"FAIL: {test.__name__}", flush=True)
            traceback.print_exc()
        else:
            print(f"PASS: {test.__name__}")
    print(f"\n{len(tests) - failures}/{len(tests)} tests passed.")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(run_tests())

